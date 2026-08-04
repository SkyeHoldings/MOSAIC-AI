# -*- coding: utf-8 -*-
import pathlib
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.enum.chart import XL_CHART_TYPE
from pptx.chart.data import CategoryChartData
from pptx.oxml.ns import qn
from lxml import etree

amp = chr(38)
dollar = chr(36)
SRC = pathlib.Path.home() / "OneDrive" / "Desktop" / "Clients" / "Surface" / ("Data Pulls " + amp + " Reporting") / "Red Robin _ Ad Schedule Analysis.pptx"
ICON = pathlib.Path(r"C:/Users/skyes/Projects/understory-marketing-site/_pptx_assets/icon.png")

ACCENT = RGBColor(0xF8, 0x30, 0x00)
BODY = RGBColor(0x16, 0x16, 0x16)
MUTED = RGBColor(0x6E, 0x6E, 0x6E)
MUTED2 = RGBColor(0x9A, 0x9A, 0x9A)
GOOD = RGBColor(0x2D, 0x86, 0x59)
BAD = RGBColor(0xC0, 0x39, 0x2B)
PANEL = RGBColor(0xF2, 0xF2, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEADER_BG = RGBColor(0x1A, 0x1A, 0x1A)
BLACK = RGBColor(0x00, 0x00, 0x00)

def set_run(run, text, font_name, size_pt, bold, color):
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color

def add_textbox(slide, left, top, width, height, paragraphs):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for item in paragraphs:
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        align = None
        runs = item
        if isinstance(item, dict):
            align = item.get("align")
            runs = item["runs"]
        if align:
            p.alignment = align
        for text, font, size, bold, color in runs:
            r = p.add_run()
            set_run(r, text, font, size, bold, color)
    return box

def solid_shape(slide, left, top, width, height, color):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    return sh

def set_cell(cell, text, font, size, bold, color, fill, align=PP_ALIGN.LEFT):
    cell.text = ""
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    set_run(r, text, font, size, bold, color)
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == qn("a:solidFill"):
            tcPr.remove(child)
    solid = etree.SubElement(tcPr, qn("a:solidFill"))
    srgb = etree.SubElement(solid, qn("a:srgbClr"))
    srgb.set("val", str(fill))
    tcPr.set("anchor", "ctr")

def find_page_number_shape(s):
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if len(t) == 2 and t.isdigit():
            return sh
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        if sh.left > 10000000 and sh.top > 6400000:
            return sh
    return None

def main():
    prs = Presentation(str(SRC))
    before = len(prs.slides)
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    sldIdLst = prs.slides._sldIdLst
    new_id = list(sldIdLst)[-1]
    sldIdLst.remove(new_id)
    sldIdLst.insert(4, new_id)
    solid_shape(slide, 0, 0, prs.slide_width, prs.slide_height, BODY)
    slide.shapes.add_picture(str(ICON), 11475720, 274320, 310896, 338328)
    solid_shape(slide, 502920, 6547104, 11183112, 11887, ACCENT)
    footer_right = "   /   RED ROBIN — AD SCHEDULE ANALYSIS"
    add_textbox(slide, 457200, 6583680, 8229600, 274320, [{"runs": [("SURFACE", "Arial Narrow", 10, True, ACCENT), (footer_right, "Arial Narrow", 10, False, MUTED2)]}])
    add_textbox(slide, 10789920, 6583680, 914400, 274320, [{"runs": [("05", "Arial Narrow", 10, False, MUTED2)], "align": PP_ALIGN.RIGHT}])
    add_textbox(slide, 502920, 274320, 10515600, 274320, [{"runs": [("PERFORMANCE MAX", "Arial Narrow", 12.5, True, ACCENT)]}])
    headline = "WEDNESDAY IS THE MOST EXPENSIVE DAY — SATURDAY THE CHEAPEST"
    add_textbox(slide, 502920, 502920, 10972800, 548640, [{"runs": [(headline, "Arial Narrow", 22, True, WHITE)]}])
    solid_shape(slide, 530352, 1051560, 658368, 68580, ACCENT)
    sub = "Average PMax CPA " + dollar + "0.39 · May 27 – Jul 27, 2026"
    add_textbox(slide, 502920, 1143000, 10972800, 274320, [{"runs": [(sub, "Arial", 12, False, MUTED2)]}])
    chart_data = CategoryChartData()
    chart_data.categories = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
    values = (0.42, 0.41, 0.44, 0.39, 0.37, 0.35, 0.39)
    chart_data.add_series("CPA (" + dollar + ")", values)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 365760, 1463040, 6400800, 3200400, chart_data)
    chart = gf.chart
    chart.has_legend = False
    chart.has_title = True
    chart.chart_title.text_frame.paragraphs[0].clear()
    tr = chart.chart_title.text_frame.paragraphs[0].add_run()
    set_run(tr, "PMax CPA by Day of Week", "Arial Narrow", 12, True, WHITE)
    try:
        chart.category_axis.has_title = True
        cat_title = chart.category_axis.axis_title
        cat_title.text_frame.paragraphs[0].clear()
        r = cat_title.text_frame.paragraphs[0].add_run()
        set_run(r, "Day of Week", "Arial", 10, False, MUTED2)
    except Exception as e:
        print("cat axis title err", e)
    try:
        chart.value_axis.has_title = True
        val_title = chart.value_axis.axis_title
        val_title.text_frame.paragraphs[0].clear()
        r = val_title.text_frame.paragraphs[0].add_run()
        set_run(r, "CPA (" + dollar + ")", "Arial", 10, False, MUTED2)
    except Exception as e:
        print("val axis title err", e)
    series = chart.series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = ACCENT
    chart.value_axis.minimum_scale = 0.30
    chart.value_axis.maximum_scale = 0.50
    chart.value_axis.has_major_gridlines = True
    avg_label = "Avg CPA " + dollar + "0.39"
    add_textbox(slide, 365760, 4622800, 4000000, 228600, [{"runs": [(avg_label, "Arial Narrow", 11, True, ACCENT), ("  ·  reference across days", "Arial", 10, False, MUTED2)]}])
    rows_data = [
        ("Day", "CPA", "vs Avg", "Cost %", "Conv %", "Efficiency"),
        ("Mon", dollar + "0.42", "+7.7%", "14.3%", "13.3%", "0.93"),
        ("Tue", dollar + "0.41", "+5.1%", "12.8%", "12.2%", "0.95"),
        ("Wed", dollar + "0.44", "+12.8%", "14.6%", "13.0%", "0.90"),
        ("Thu", dollar + "0.39", "0%", "14.7%", "14.8%", "1.01"),
        ("Fri", dollar + "0.37", "-5.1%", "15.1%", "16.0%", "1.06"),
        ("Sat", dollar + "0.35", "-10.3%", "14.5%", "16.5%", "1.14"),
        ("Sun", dollar + "0.39", "0%", "14.0%", "14.2%", "1.01"),
    ]
    vs_colors = {3: BAD, 6: GOOD}
    tbl_shape = slide.shapes.add_table(8, 6, 6900000, 1463040, 4800600, 2926080)
    table = tbl_shape.table
    widths = [700000, 800000, 900000, 800000, 800000, 900000]
    for i, w in enumerate(widths):
        table.columns[i].width = w
    for ri, row in enumerate(rows_data):
        for ci, val in enumerate(row):
            cell = table.cell(ri, ci)
            if ri == 0:
                set_cell(cell, val, "Arial Narrow", 9, True, WHITE, HEADER_BG, PP_ALIGN.CENTER)
            else:
                color = BLACK
                if ci == 2:
                    color = vs_colors.get(ri, MUTED)
                align = PP_ALIGN.CENTER if ci > 0 else PP_ALIGN.LEFT
                set_cell(cell, val, "Arial", 9, False, color, WHITE, align)
    solid_shape(slide, 502920, 5074920, 11064240, 1143000, PANEL)
    solid_shape(slide, 502920, 5074920, 109728, 1143000, ACCENT)
    callout = (
        "Spend is fairly even across days (~13–15%). Efficiency gap is CPA-driven — "
        "Wednesday overpays vs conversion share; Saturday underpays. Still not a blackout "
        "lever (can" + chr(39) + "t cut whole days without hour blackouts), but useful for pacing context."
    )
    add_textbox(slide, 758952, 5166360, 10698480, 960120, [
        {"runs": [("CALLOUT", "Arial Narrow", 11, True, ACCENT)]},
        {"runs": [(callout, "Arial", 12, False, MUTED)]},
    ])
    for idx, s in enumerate(prs.slides, 1):
        sh = find_page_number_shape(s)
        if sh is None:
            print("NO PAGE NUM on slide", idx)
            continue
        tf = sh.text_frame
        p = tf.paragraphs[0]
        font_name, size, bold, color = "Arial Narrow", 10, False, MUTED2
        if p.runs:
            r0 = p.runs[0]
            if r0.font.name:
                font_name = r0.font.name
            if r0.font.size:
                size = r0.font.size.pt
            if r0.font.bold is not None:
                bold = r0.font.bold
            try:
                if r0.font.color.rgb:
                    color = r0.font.color.rgb
            except Exception:
                pass
        p.clear()
        r = p.add_run()
        set_run(r, f"{idx:02d}", font_name, size, bold, color)
        print(f"set page {idx:02d}")
    prs.save(str(SRC))
    print(f"SAVED. before={before} after={len(prs.slides)}")

if __name__ == "__main__":
    main()
