import json
import re
from openpyxl import load_workbook

path = r"C:\Users\skyes\Projects\understory-marketing-site\contact-list-apollo-copy.xlsx"
out_path = r"C:\Users\skyes\Projects\understory-marketing-site\contact-list-apollo-dump.txt"

wb = load_workbook(path, data_only=True)

lines = []
lines.append("=== SHEET NAMES ===")
for name in wb.sheetnames:
    lines.append(repr(name))

need_pattern = re.compile(r"need\s*enriched", re.I)
enriched_pattern = re.compile(r"^enriched$", re.I)

need_sheet = None
enriched_sheet = None

for name in wb.sheetnames:
    if need_pattern.search(name):
        need_sheet = name
    if enriched_pattern.match(name.strip()):
        enriched_sheet = name

lines.append("")
lines.append("=== SHEET METADATA ===")

for name in wb.sheetnames:
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    max_row = len(rows)
    max_col = max((len(r) for r in rows), default=0)
    headers = list(rows[0]) if rows else []
    header_vals = [("" if h is None else str(h)) for h in headers]
    lines.append(f"Sheet: {name!r}")
    lines.append(f"  max_row={max_row}, max_column={max_col}")
    lines.append(f"  headers ({len(header_vals)}): {header_vals}")

if need_sheet is None:
    for name in wb.sheetnames:
        if "enrich" in name.lower() and "need" in name.lower():
            need_sheet = name
            break

lines.append("")
lines.append(f"=== NEED ENRICHED SHEET: {need_sheet!r} ===")

if need_sheet:
    ws = wb[need_sheet]
    all_rows = list(ws.iter_rows(values_only=True))
    if all_rows:
        headers = [("" if h is None else str(h)) for h in all_rows[0]]
        lines.append("FORMAT: JSONL (one object per row)")
        for row in all_rows[1:]:
            obj = {}
            for j, h in enumerate(headers):
                key = h if h else f"col_{j+1}"
                val = row[j] if j < len(row) else None
                if val is not None and not isinstance(val, (str, int, float, bool)):
                    val = str(val)
                obj[key] = val
            lines.append(json.dumps(obj, ensure_ascii=False, default=str))
    else:
        lines.append("(empty sheet)")
else:
    lines.append("NOT FOUND. Available sheets: " + ", ".join(wb.sheetnames))

if enriched_sheet is None:
    for name in wb.sheetnames:
        if name.strip().lower() == "enriched":
            enriched_sheet = name
            break

lines.append("")
lines.append(f"=== ENRICHED SHEET SAMPLE: {enriched_sheet!r} ===")

if enriched_sheet:
    ws = wb[enriched_sheet]
    rows = list(ws.iter_rows(values_only=True))
    if rows:
        headers = [("" if h is None else str(h)) for h in rows[0]]
        lines.append(f"HEADERS ({len(headers)}): {json.dumps(headers, ensure_ascii=False)}")
        for idx, row in enumerate(rows[1:4], start=1):
            obj = {}
            for j, h in enumerate(headers):
                key = h if h else f"col_{j+1}"
                val = row[j] if j < len(row) else None
                if val is not None and not isinstance(val, (str, int, float, bool)):
                    val = str(val)
                obj[key] = val
            lines.append(f"DATA_ROW_{idx}: {json.dumps(obj, ensure_ascii=False, default=str)}")
    else:
        lines.append("(empty sheet)")
else:
    lines.append("NOT FOUND")

wb.close()

with open(out_path, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))

print(f"Wrote {len(lines)} lines to {out_path}")
