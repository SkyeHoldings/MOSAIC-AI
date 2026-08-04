"""
Reveal Apollo direct dials via webhook.site and write them into the Excel file.
Requires APOLLO_API_KEY and WEBHOOK_UUID env vars.
"""
from __future__ import annotations

import json
import os
import re
import shutil
import time
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

import openpyxl

API_KEY = os.environ["APOLLO_API_KEY"].strip()
WEBHOOK_UUID = os.environ["WEBHOOK_UUID"].strip()
WEBHOOK_URL = f"https://webhook.site/{WEBHOOK_UUID}"

SRC_CANDIDATES = [
    Path(r"c:\Users\skyes\OneDrive\Desktop\Contact List for Apollo - Direct Dials.xlsx"),
    Path(r"c:\Users\skyes\OneDrive\Desktop\Contact List for Apollo - Enriched - UPDATED.xlsx"),
    Path(r"c:\Users\skyes\OneDrive\Desktop\Contact List for Apollo - Enriched.xlsx"),
    Path(r"C:\Users\skyes\Projects\understory-marketing-site\_tmp_apollo_enrich_work.xlsx"),
]
SKIP_IF_STATUS_CONTAINS = "direct_dial"
WORK = Path(r"C:\Users\skyes\Projects\understory-marketing-site\_tmp_apollo_direct_work.xlsx")
OUT = Path(r"c:\Users\skyes\OneDrive\Desktop\Contact List for Apollo - Direct Dials.xlsx")
LOG = Path(r"C:\Users\skyes\Projects\understory-marketing-site\_tmp_apollo_direct_log.json")
MAP = Path(r"C:\Users\skyes\Projects\understory-marketing-site\_tmp_apollo_person_row_map.json")

HEADERS = {
    "X-Api-Key": API_KEY,
    "Content-Type": "application/json",
    "Cache-Control": "no-cache",
    "accept": "application/json",
}
EMPTY = {"", "none", "-", "n/a", "na", "null", "?"}


def blank(v) -> bool:
    return v is None or str(v).strip().lower() in EMPTY


def looks_like_company_mainline(phone: str | None) -> bool:
    """Heuristic: many HQ/main lines we previously filled; still try reveal for everyone missing a mobile-looking personal dial."""
    if blank(phone):
        return True
    return True  # always attempt reveal for rows we target; overwrite with direct dial when found


def api(method: str, url: str, payload=None, query=None):
    if query:
        url += ("&" if "?" in url else "?") + urllib.parse.urlencode(query, doseq=True)
    data = None if payload is None else json.dumps(payload).encode()
    req = urllib.request.Request(url, data=data, headers=HEADERS, method=method)
    try:
        with urllib.request.urlopen(req, timeout=90) as resp:
            return resp.status, json.loads(resp.read().decode() or "{}")
    except urllib.error.HTTPError as e:
        body = e.read().decode(errors="replace")
        try:
            parsed = json.loads(body)
        except Exception:
            parsed = {"raw": body[:2000]}
        return e.code, parsed


def webhook_get(path: str, query=None):
    url = f"https://webhook.site/token/{WEBHOOK_UUID}{path}"
    if query:
        url += "?" + urllib.parse.urlencode(query, doseq=True)
    req = urllib.request.Request(url, headers={"Accept": "application/json"})
    with urllib.request.urlopen(req, timeout=60) as resp:
        return json.loads(resp.read().decode() or "{}")


def people_match(email=None, name=None, organization_name=None, person_id=None):
    payload = {
        "reveal_phone_number": True,
        "reveal_personal_emails": False,
        "run_waterfall_phone": True,
        "webhook_url": WEBHOOK_URL,
    }
    if person_id:
        payload["id"] = person_id
    if email:
        payload["email"] = email
    if name:
        payload["name"] = name
    if organization_name:
        payload["organization_name"] = organization_name
    q = {
        "reveal_phone_number": "true",
        "run_waterfall_phone": "true",
        "webhook_url": WEBHOOK_URL,
    }
    return api("POST", "https://api.apollo.io/api/v1/people/match", payload, q)


def bulk_match(details: list[dict]):
    payload = {
        "reveal_phone_number": True,
        "reveal_personal_emails": False,
        "run_waterfall_phone": True,
        "webhook_url": WEBHOOK_URL,
        "details": details,
    }
    q = {
        "reveal_phone_number": "true",
        "run_waterfall_phone": "true",
        "webhook_url": WEBHOOK_URL,
    }
    return api("POST", "https://api.apollo.io/api/v1/people/bulk_match", payload, q)


def prefer_direct(phones: list[dict]) -> str | None:
    if not phones:
        return None
    order = {"mobile": 0, "work_direct": 1, "home": 2, "other": 3, "hq": 4}
    ranked = sorted(
        phones,
        key=lambda p: (
            order.get((p.get("type_cd") or "").lower(), 9),
            0 if (p.get("status_cd") or "") == "valid_number" else 1,
        ),
    )
    top = ranked[0]
    return top.get("sanitized_number") or top.get("raw_number")


def clean_name(raw) -> str | None:
    if blank(raw):
        return None
    s = str(raw)
    s = re.split(r"[-–—|/]", s)[0]
    s = re.sub(
        r"\b(owner|bm|abm|liaison|compliance admin|branch manager|area branch manager|office manager|gm|manager)\b",
        "",
        s,
        flags=re.I,
    )
    s = re.sub(r"[?]+", "", s).strip(" ,")
    if "&" in s or "," in s:
        s = re.split(r"[&,]", s)[0].strip()
    return s if len(s) >= 2 else None


def pick_src() -> Path:
    for p in SRC_CANDIDATES:
        if p.exists():
            return p
    raise SystemExit("No source workbook found")


def main():
    src = pick_src()
    print("SRC", src)
    print("WEBHOOK", WEBHOOK_URL)
    shutil.copy2(src, WORK)
    wb = openpyxl.load_workbook(WORK)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}
    col = {k: idx[k] + 1 for k in [
        "Business", "Contact", "Email / Contact Info", "Apollo Email", "Phone",
        "Apollo Name", "Apollo Title", "LinkedIn", "Enrichment Status",
    ]}

    # Build targets: rows with an email (best for match) OR apollo name + business
    targets = []
    for row in range(2, ws.max_row + 1):
        business = ws.cell(row, col["Business"]).value
        contact = ws.cell(row, col["Contact"]).value
        email = ws.cell(row, col["Email / Contact Info"]).value or ws.cell(row, col["Apollo Email"]).value
        phone = ws.cell(row, col["Phone"]).value
        apollo_name = ws.cell(row, col["Apollo Name"]).value
        status = str(ws.cell(row, col["Enrichment Status"]).value or "")
        if blank(business) and blank(contact) and blank(email):
            continue
        if SKIP_IF_STATUS_CONTAINS in status.lower():
            continue  # already has a direct dial from prior pass
        # Prefer enriching everyone we can identify; focus rows that have email or apollo name
        detail = {"row": row, "business": business, "contact": contact, "email": None, "name": None}
        if email and "@" in str(email) and " " not in str(email).split("@")[0]:
            detail["email"] = str(email).strip()
        name = clean_name(apollo_name) or clean_name(contact)
        if name:
            detail["name"] = name
        if not detail["email"] and not detail["name"]:
            continue
        targets.append(detail)

    print(f"targets={len(targets)}")

    person_to_rows: dict[str, list[int]] = {}
    row_meta = {}
    requested = 0
    errors = []

    # Request in chunks of 10 via bulk_match where possible
    i = 0
    while i < len(targets):
        chunk = targets[i : i + 10]
        details = []
        for t in chunk:
            d = {}
            if t["email"]:
                d["email"] = t["email"]
            if t["name"]:
                d["name"] = t["name"]
            if t["business"]:
                d["organization_name"] = str(t["business"])
            if d:
                details.append(d)
                row_meta[t["row"]] = t
        if not details:
            i += 10
            continue

        st, data = bulk_match(details)
        print(f"bulk[{i}:{i+len(details)}] status={st}")
        if st >= 400:
            errors.append({"chunk": i, "status": st, "error": data})
            # fallback single matches
            for t in chunk:
                payload = {}
                if t["email"]:
                    st2, data2 = people_match(email=t["email"])
                elif t["name"]:
                    st2, data2 = people_match(name=t["name"], organization_name=t.get("business"))
                else:
                    continue
                person = (data2 or {}).get("person") or {}
                pid = person.get("id")
                if pid:
                    person_to_rows.setdefault(pid, []).append(t["row"])
                    requested += 1
                    print(f"  single row {t['row']} -> {pid} {person.get('name')}")
                else:
                    print(f"  single miss row {t['row']} st={st2}")
                time.sleep(0.25)
            i += 10
            time.sleep(0.5)
            continue

        matches = data.get("matches") or data.get("people") or []
        # bulk_match returns {matches:[{status, person:{...}}]} typically
        if matches and isinstance(matches[0], dict) and "person" in matches[0]:
            people = [m.get("person") for m in matches]
        else:
            people = matches

        for t, person in zip(chunk, people + [None] * len(chunk)):
            if not person:
                # try individual
                if t["email"]:
                    st2, data2 = people_match(email=t["email"])
                    person = (data2 or {}).get("person")
                elif t["name"]:
                    st2, data2 = people_match(name=t["name"], organization_name=t.get("business"))
                    person = (data2 or {}).get("person")
            if not person:
                print(f"  no person row {t['row']}")
                continue
            pid = person.get("id")
            if not pid:
                continue
            person_to_rows.setdefault(pid, []).append(t["row"])
            requested += 1
            print(f"  queued row {t['row']} -> {pid} {person.get('name')}")
        i += 10
        time.sleep(0.6)

    MAP.write_text(json.dumps({"person_to_rows": person_to_rows, "requested": requested}, indent=2), encoding="utf-8")
    print(f"queued_person_ids={len(person_to_rows)} requested={requested}")

    # Poll webhook for phone payloads (paginate so we don't miss callbacks)
    found: dict[str, str] = {}
    deadline = time.time() + 60 * 20  # up to 20 minutes
    seen_request_ids = set()
    print("polling webhook for direct dials...")
    while time.time() < deadline and len(found) < len(person_to_rows):
        try:
            page_num = 1
            while page_num <= 30:
                page = webhook_get(
                    "/requests",
                    {"sorting": "newest", "per_page": 50, "page": page_num},
                )
                data = page.get("data") or []
                if not data:
                    break
                for item in data:
                    rid = item.get("uuid") or item.get("id")
                    if rid in seen_request_ids:
                        continue
                    seen_request_ids.add(rid)
                    content = item.get("content") or item.get("request") or ""
                    if isinstance(content, dict):
                        payload = content
                    else:
                        try:
                            payload = json.loads(content or "{}")
                        except Exception:
                            continue
                    people = payload.get("people") or []
                    for p in people:
                        pid = p.get("id") or p.get("_id")
                        phone = prefer_direct(p.get("phone_numbers") or [])
                        if pid and phone and pid not in found:
                            found[pid] = phone
                            print(
                                f"  dial {pid} -> {phone} ({(p.get('phone_numbers') or [{}])[0].get('type_cd')})"
                            )
                if len(data) < 50:
                    break
                page_num += 1
        except Exception as e:
            print("poll error", e)
            time.sleep(3)
            continue
        print(f"  progress {len(found)}/{len(person_to_rows)}")
        if len(found) >= len(person_to_rows):
            break
        time.sleep(5)

    # Apply to sheet
    filled = 0
    for pid, rows in person_to_rows.items():
        phone = found.get(pid)
        if not phone:
            continue
        for row in rows:
            ws.cell(row, col["Phone"]).value = phone
            old = ws.cell(row, col["Enrichment Status"]).value or ""
            ws.cell(row, col["Enrichment Status"]).value = f"{old}; direct_dial".strip("; ")[:240]
            filled += 1

    wb.save(WORK)
    try:
        shutil.copy2(WORK, OUT)
        saved = str(OUT)
    except PermissionError:
        alt = OUT.with_name("Contact List for Apollo - Direct Dials - COPY.xlsx")
        shutil.copy2(WORK, alt)
        saved = str(alt)

    summary = {
        "targets": len(targets),
        "queued_people": len(person_to_rows),
        "dials_received": len(found),
        "rows_updated": filled,
        "errors": errors[:20],
        "saved": saved,
        "found": found,
    }
    LOG.write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print("DONE", json.dumps({k: summary[k] for k in summary if k != "found"}, indent=2))


if __name__ == "__main__":
    main()
