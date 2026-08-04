import openpyxl
import re
import json

src = r"c:\Users\skyes\OneDrive\Desktop\Contact List for Apollo - Enriched.xlsx"
dst = r"C:\Users\skyes\Projects\understory-marketing-site\_tmp_apollo_contacts.xlsx"

import shutil
shutil.copy2(src, dst)

wb = openpyxl.load_workbook(dst, data_only=True)
print("SHEETS:", wb.sheetnames)

phone_pat = re.compile(r"phone|mobile|direct.?dial|work.?phone|corporate.?phone", re.I)
empty_vals = {"", "none", "-", "n/a", "na", "null"}

summary = {}

for s in wb.sheetnames:
    ws = wb[s]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        print(f"\n=== {s}: empty ===")
        continue
    headers = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(rows[0])]
    print(f"\n=== {s} ({len(rows)-1} data rows) ===")
    print("HEADERS:", headers)
    phone_cols = [i for i, h in enumerate(headers) if phone_pat.search(h or "")]
    print("PHONE COLS:", [(i, headers[i]) for i in phone_cols])
    interesting = []
    for i, h in enumerate(headers):
        if re.search(
            r"company|organization|name|title|email|website|domain|linkedin|person|first|last|job",
            h or "",
            re.I,
        ):
            interesting.append((i, h))
    print("KEY COLS:", interesting[:40])

    has_any = missing = 0
    missing_rows = []
    if phone_cols:
        for r in rows[1:]:
            vals = [r[i] if i < len(r) else None for i in phone_cols]
            if any(
                v is not None and str(v).strip().lower() not in empty_vals for v in vals
            ):
                has_any += 1
            else:
                missing += 1
                d = {
                    headers[i]: (r[i] if i < len(r) else None)
                    for i in range(len(headers))
                }
                missing_rows.append(d)
        print(f"PHONE COVERAGE: has={has_any} missing={missing}")

    for ri, r in enumerate(rows[1:3], 1):
        d = {headers[i]: (r[i] if i < len(r) else None) for i in range(len(headers))}
        print(f"SAMPLE {ri}:", {k: d[k] for k in list(d)[:20]})

    summary[s] = {
        "total": len(rows) - 1,
        "has_phone": has_any,
        "missing_phone": missing,
        "headers": headers,
        "phone_cols": [headers[i] for i in phone_cols],
        "missing_preview": missing_rows[:25],
    }

out = r"C:\Users\skyes\Projects\understory-marketing-site\_tmp_apollo_phone_gap.json"
with open(out, "w", encoding="utf-8") as f:
    json.dump(summary, f, default=str, indent=2)
print("\nWrote", out)
