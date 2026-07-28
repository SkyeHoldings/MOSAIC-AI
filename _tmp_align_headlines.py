"""Strict column alignment: one text/role per headline column."""
import re
from collections import Counter

import openpyxl

path = r"c:\Users\skyes\OneDrive\Desktop\RR Copy.xlsx"
backup = r"c:\Users\skyes\OneDrive\Desktop\RR Copy.backup-before-align.xlsx"

wb = openpyxl.load_workbook(backup)
ws = wb["Bulksheet"]
headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
h_cols = [
    i
    for i, h in enumerate(headers, 1)
    if h and str(h).startswith("Headline") and not str(h).lower().startswith("long")
]
city_re = re.compile(r"^(Red Robin Near |Visit Red Robin )", re.I)

text_count = Counter()
for r in range(2, ws.max_row + 1):
    for c in h_cols:
        v = ws.cell(r, c).value
        if v:
            text_count[str(v).strip()] += 1
shared_texts = {t for t, n in text_count.items() if n >= 1000}
entree = next(t for t in shared_texts if t.startswith("Pick Your Favorite"))

CORE_ORDER = [
    "Red Robin's Big Yummm Deals",
    "Available for a Limited Time",
    "Big Deals, Small Price",
    "Find Restaurants Near You",
    "Big Yummm. Small Price.",
    entree,
    "Bottomless Sides & Beverages",
    "Big Yummm and Bottomless Side",
    "$9.99 for a Limited Time!",
    "Visit Your Local Red Robin",
    "Budget-Friendly Bites",
    "Deals This Good Go Fast",
    "Skip the Kitchen Tonight",
]
DUAL_H14 = "$9.99 Big Yummm Deals"
DUAL_H15 = "Big Yummm for $9.99"
core_to_col = {text: h_cols[i] for i, text in enumerate(CORE_ORDER)}
col_h14, col_h15 = h_cols[13], h_cols[14]
stats = Counter()

for r in range(2, ws.max_row + 1):
    existing = [
        str(ws.cell(r, c).value).strip()
        for c in h_cols
        if ws.cell(r, c).value and str(ws.cell(r, c).value).strip()
    ]
    cities = [v for v in existing if city_re.match(v)]
    shared_set = {v for v in existing if v in shared_texts}
    near = next((v for v in cities if v.lower().startswith("red robin near")), None)
    visit = next((v for v in cities if v.lower().startswith("visit red robin")), None)

    for c in h_cols:
        ws.cell(r, c).value = None

    for text, col in core_to_col.items():
        if text in shared_set:
            ws.cell(r, col).value = text

    if near or visit:
        stats["city_rows"] += 1
        if near:
            ws.cell(r, col_h14).value = near
            if DUAL_H14 in shared_set:
                stats["removed_dual"] += 1
        elif DUAL_H14 in shared_set:
            ws.cell(r, col_h14).value = DUAL_H14
        if visit:
            ws.cell(r, col_h15).value = visit
            if DUAL_H15 in shared_set:
                stats["removed_dual"] += 1
        elif DUAL_H15 in shared_set:
            ws.cell(r, col_h15).value = DUAL_H15
    else:
        stats["noncity_rows"] += 1
        if DUAL_H14 in shared_set:
            ws.cell(r, col_h14).value = DUAL_H14
        if DUAL_H15 in shared_set:
            ws.cell(r, col_h15).value = DUAL_H15

out = openpyxl.load_workbook(path)
idx = out.sheetnames.index("Bulksheet")
del out["Bulksheet"]
new_bs = out.create_sheet("Bulksheet", idx)
for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
    for cell in row:
        new_bs.cell(cell.row, cell.column, cell.value)
out.save(path)
print("Aligned. Stats:", dict(stats))
for i, t in enumerate(CORE_ORDER, 1):
    print(f"  H{i}: {t}")
print(f"  H14: Near city  |  else {DUAL_H14}")
print(f"  H15: Visit city |  else {DUAL_H15}")
