# -*- coding: utf-8 -*-
"""Insert PMax conversions-by-conversion-time HOD slide after slide 6."""
import pathlib
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.enum.chart import XL_CHART_TYPE
from pptx.chart.data import CategoryChartData
from pptx.oxml.ns import qn
from lxml import etree

amp = chr(38)
dollar = chr(36)
en = "\u2013"
mdash = "\u2014"
mid = "\u00b7"
star = "*"

SRC = pathlib.Path.home() / "OneDrive" / "Desktop" / "Clients" / "Surface" / ("Data Pulls " + amp + " Reporting") / "Red Robin _ Ad Schedule Analysis.pptx"
ICON = pathlib.Path(r"C:/Users/skyes/Projects/understory-marketing-site/_pptx_assets/icon.png")

ACCENT = RGBColor(0xF8, 0x30, 0x00)
BODY = RGBColor(0x16, 0x16, 0x16)
MUTED = RGBColor(0x6E, 0x6E, 0x6E)
MUTED2 = RGBColor(0x9A, 0x9A, 0x9A)
GOOD = RGBColor(0x2D, 0x86, 0x59)
BAD = RGBColor(0xC0, 0x39, 0x2B)
PANEL = RGBColor(0xF2, 0xF2, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEADER_BG = RGBColor(0x1A, 0x1A, 0x1A)
BLACK = RGBColor(0x00, 0x00, 0x00)


def set_run(run, text, font_name, size_pt, bold, color):
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color


def add_textbox(slide, left, top, width, height, paragraphs):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for item in paragraphs:
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        align = None
        runs = item
        if isinstance(item, dict):
            align = item.get("align")
            runs = item["runs"]
        if align:
            p.alignment = align
        for text, font, size, bold, color in runs:
            r = p.add_run()
            set_run(r, text, font, size, bold, color)
    return box


def solid_shape(slide, left, top, width, height, color):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    sh.line.fill.background()
    return sh


def set_cell(cell, text, font, size, bold, color, fill, align=PP_ALIGN.LEFT):
    cell.text = ""
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    set_run(r, text, font, size, bold, color)
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == qn("a:solidFill"):
            tcPr.remove(child)
    solid = etree.SubElement(tcPr, qn("a:solidFill"))
    srgb = etree.SubElement(solid, qn("a:srgbClr"))
    srgb.set("val", str(fill))
    tcPr.set("anchor", "ctr")


def find_page_number_shape(s):
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip()
        if len(t) == 2 and t.isdigit():
            return sh
    for sh in s.shapes:
        if not sh.has_text_frame:
            continue
        if sh.left > 10000000 and sh.top > 6400000:
            return sh
    return None


def style_category_axis_font(chart, size_pt=8):
    try:
        cat = chart.category_axis
        txPr = cat._element.find(qn("c:txPr"))
        if txPr is None:
            from pptx.oxml import parse_xml
            txPr = parse_xml(
                '<c:txPr xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
                'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
                "<a:bodyPr/><a:lstStyle/>"
                '<a:p><a:pPr><a:defRPr sz="%d"/></a:pPr><a:endParaRPr/></a:p>'
                "</c:txPr>" % int(size_pt * 100)
            )
            cat._element.append(txPr)
        else:
            defRPr = txPr.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}defRPr")
            if defRPr is not None:
                defRPr.set("sz", str(int(size_pt * 100)))
    except Exception as e:
        print("cat font err", e)


def main():
    prs = Presentation(str(SRC))
    before = len(prs.slides)
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    sldIdLst = prs.slides._sldIdLst
    new_id = list(sldIdLst)[-1]
    sldIdLst.remove(new_id)
    # After current slide 6 (0-based index 5) => insert at position 6 => new slide 7
    sldIdLst.insert(6, new_id)

    solid_shape(slide, 0, 0, prs.slide_width, prs.slide_height, BODY)
    slide.shapes.add_picture(str(ICON), 11475720, 274320, 310896, 338328)
    solid_shape(slide, 502920, 6547104, 11183112, 11887, ACCENT)

    footer_right = "   /   RED ROBIN " + mdash + " AD SCHEDULE ANALYSIS"
    add_textbox(
        slide, 457200, 6583680, 8229600, 274320,
        [{"runs": [("SURFACE", "Arial Narrow", 10, True, ACCENT), (footer_right, "Arial Narrow", 10, False, MUTED2)]}],
    )
    add_textbox(
        slide, 10789920, 6583680, 914400, 274320,
        [{"runs": [("07", "Arial Narrow", 10, False, MUTED2)], "align": PP_ALIGN.RIGHT}],
    )

    add_textbox(
        slide, 502920, 274320, 10515600, 274320,
        [{"runs": [("PERFORMANCE MAX", "Arial Narrow", 12.5, True, ACCENT)]}],
    )
    headline = "CONVERSIONS BY TIME PEAK 1P" + en + "6P " + mdash + " SAME AFTERNOON WINDOW AS SPEND"
    add_textbox(
        slide, 502920, 502920, 10972800, 548640,
        [{"runs": [(headline, "Arial Narrow", 22, True, WHITE)]}],
    )
    solid_shape(slide, 530352, 1051560, 658368, 68580, ACCENT)
    sub = (
        "Conversions by conversion time " + mid + " Jun 22 " + en + " Jul 27, 2026 "
        + mid + " Avg CPA " + dollar + "1.57"
    )
    add_textbox(
        slide, 502920, 1143000, 10972800, 274320,
        [{"runs": [(sub, "Arial", 12, False, MUTED2)]}],
    )

    chart_data = CategoryChartData()
    chart_data.categories = [str(h) for h in range(24)]
    values = (
        482.5, 298.0, 150.0, 104.0, 92.0, 83.0, 728.0, 1218.0,
        2009.0, 3204.67, 6200.02, 20648.0, 37917.62, 42634.06, 45749.17, 44378.06,
        42378.8, 44303.81, 43394.86, 36916.85, 28053.25, 16563.04, 8129.49, 1302.93,
    )
    chart_data.add_series("Conversions", values)
    gf = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        274320, 1463040, 7000000, 3200400,
        chart_data,
    )
    chart = gf.chart
    chart.has_legend = False
    chart.has_title = True
    chart.chart_title.text_frame.paragraphs[0].clear()
    tr = chart.chart_title.text_frame.paragraphs[0].add_run()
    set_run(tr, "PMax Conversions by Hour (Conversion Time)", "Arial Narrow", 12, True, WHITE)

    try:
        chart.category_axis.has_title = True
        cat_title = chart.category_axis.axis_title
        cat_title.text_frame.paragraphs[0].clear()
        r = cat_title.text_frame.paragraphs[0].add_run()
        set_run(r, "Hour of Day", "Arial", 9, False, MUTED2)
    except Exception as e:
        print("cat axis title err", e)

    try:
        chart.value_axis.has_title = True
        val_title = chart.value_axis.axis_title
        val_title.text_frame.paragraphs[0].clear()
        r = val_title.text_frame.paragraphs[0].add_run()
        set_run(r, "Conversions", "Arial", 9, False, MUTED2)
    except Exception as e:
        print("val axis title err", e)

    series = chart.series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = ACCENT
    chart.value_axis.minimum_scale = 0
    chart.value_axis.maximum_scale = 50000
    chart.value_axis.has_major_gridlines = True
    style_category_axis_font(chart, 7)

    note = (
        star + "CPA = click-time cost / conv-time conversions (interpret carefully). "
        "Hours 0" + en + "5: conversions with " + dollar + "0 click cost " + mdash + " lag artifact."
    )
    add_textbox(
        slide, 274320, 4622800, 7000000, 320000,
        [{"runs": [(note, "Arial", 9, False, MUTED2)]}],
    )

    rows_data = [
        ("Hour band", "Conv", "Conv %", "Cost %", "CPA" + star, "Read"),
        (
            "12a" + en + "5a (0" + en + "5)",
            "1,210",
            "0.28%",
            "0.0%",
            "n/a",
            "Conv-time only; " + dollar + "0 click cost " + mdash + " lag artifact",
        ),
        (
            "6a" + en + "10a (6" + en + "10)",
            "13,360",
            "3.1%",
            "6.6%",
            dollar + "3.03" + en + "4.11",
            "Worst CPA; low volume",
        ),
        (
            "11a" + en + "12p (11" + en + "12)",
            "58,566",
            "13.7%",
            "14.5%",
            dollar + "1.54" + en + "1.88",
            "Ramping into peak",
        ),
        (
            "1p" + en + "6p (13" + en + "18)",
            "262,839",
            "61.6%",
            "59.7%",
            dollar + "1.46" + en + "1.55",
            "Peak conversion window",
        ),
        (
            "7p" + en + "11p (19" + en + "23)",
            "90,966",
            "21.3%",
            "19.1%",
            dollar + "0.54" + en + "1.46",
            "Strong; late hours cheapest CPA",
        ),
    ]

    # CPA column tint: 1 muted (n/a), 2 bad, 3 muted, 4 good, 5 good
    cpa_colors = {1: MUTED, 2: BAD, 3: MUTED, 4: GOOD, 5: GOOD}

    tbl_shape = slide.shapes.add_table(6, 6, 7380000, 1463040, 4500000, 3000000)
    table = tbl_shape.table
    widths = [980000, 620000, 560000, 560000, 720000, 1060000]
    total_w = sum(widths)
    target = 4500000
    widths = [int(w * target / total_w) for w in widths]
    for i, w in enumerate(widths):
        table.columns[i].width = w

    for ri, row in enumerate(rows_data):
        for ci, val in enumerate(row):
            cell = table.cell(ri, ci)
            if ri == 0:
                set_cell(cell, val, "Arial Narrow", 8, True, WHITE, HEADER_BG, PP_ALIGN.CENTER)
            else:
                color = BLACK
                if ci == 4:
                    color = cpa_colors.get(ri, MUTED)
                if ci == 0 or ci == 5:
                    align = PP_ALIGN.LEFT
                    size = 7.5
                else:
                    align = PP_ALIGN.CENTER
                    size = 8
                set_cell(cell, val, "Arial", size, False, color, WHITE, align)

    solid_shape(slide, 502920, 5074920, 11064240, 1143000, PANEL)
    solid_shape(slide, 502920, 5074920, 109728, 1143000, ACCENT)
    callout = (
        "By conversion time, ~62% of PMax conversions land 1p" + en + "6p "
        + mdash + " aligning with peak spend. Morning (6a" + en + "10a) is inefficient on this view "
        "(" + dollar + "3" + en + "4 CPA) but small volume. Overnight conversions without cost reflect "
        "store-visit/conversion lag, not free overnight demand."
    )
    add_textbox(
        slide, 758952, 5166360, 10698480, 960120,
        [
            {"runs": [("CALLOUT", "Arial Narrow", 11, True, ACCENT)]},
            {"runs": [(callout, "Arial", 12, False, MUTED)]},
        ],
    )

    for idx, s in enumerate(prs.slides, 1):
        sh = find_page_number_shape(s)
        if sh is None:
            print("NO PAGE NUM on slide", idx)
            continue
        tf = sh.text_frame
        p = tf.paragraphs[0]
        font_name, size, bold, color = "Arial Narrow", 10, False, MUTED2
        if p.runs:
            r0 = p.runs[0]
            if r0.font.name:
                font_name = r0.font.name
            if r0.font.size:
                size = r0.font.size.pt
            if r0.font.bold is not None:
                bold = r0.font.bold
            try:
                if r0.font.color.rgb:
                    color = r0.font.color.rgb
            except Exception:
                pass
        p.clear()
        r = p.add_run()
        set_run(r, f"{idx:02d}", font_name, size, bold, color)
        print(f"set page {idx:02d}")

    prs.save(str(SRC))
    after = len(prs.slides)
    print(f"SAVED. before={before} after={after}")

    prs2 = Presentation(str(SRC))
    s7 = prs2.slides[6]
    has_chart = any(sh.has_chart for sh in s7.shapes)
    texts = []
    for sh in s7.shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if t:
                texts.append(t[:120].replace("\n", " | "))
    print("new_slide_index", 7)
    print("slide7_has_chart", has_chart)
    print("slide_count", len(prs2.slides))
    print("slide7_texts:")
    for t in texts[:8]:
        print(" ", t.encode("ascii", "replace").decode("ascii"))

    # Confirm chart series values briefly
    for sh in s7.shapes:
        if sh.has_chart:
            ch = sh.chart
            print("chart_title", ch.chart_title.text_frame.text.encode("ascii", "replace").decode("ascii"))
            print("chart_type", ch.chart_type)
            print("series_count", len(ch.series))
            break


if __name__ == "__main__":
    main()
