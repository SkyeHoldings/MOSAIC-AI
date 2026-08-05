"""Rank companies easiest (1) -> hardest by likely ad-budget / decision difficulty."""
import re
import openpyxl

path = r"c:\Users\skyes\OneDrive\Desktop\Contact List for Apollo - Direct Dials.xlsx"
wb = openpyxl.load_workbook(path)

# Lower score = easier (smaller / lower ad budget / faster decisions)
INDUSTRY_BASE = {
    "Beauty, Hair, Nails & Massage": 10,
    "Tattoo & Body Art": 12,
    "Med Spa, Aesthetics & IV Wellness": 18,
    "Nutrition & Weight Management": 20,
    "Maternity, Doula & Family Support": 22,
    "Holistic, Naturopathic & Functional Medicine": 25,
    "Restaurants, Food & Beverage": 28,
    "Fitness & Athletic Training": 30,
    "Retail & Consumer Goods": 32,
    "Dentistry": 38,
    "Orthodontics": 40,
    "Physical Therapy & Sports Medicine": 42,
    "Mental & Behavioral Health": 45,
    "Prosthetics & Adaptive Medical": 48,
    "Construction, Trades & Home Services": 50,
    "Real Estate Brokerages": 52,
    "Legal": 55,
    "Engineering": 58,
    "Engineering & Professional Services": 58,
    "Home Care, Hospice & Senior Living": 60,
    "Mortgage & Lending": 65,
    "Title & Escrow": 68,
    "Insurance & Financial Services": 70,
    "Apartments & Housing": 72,
    "Recreation, Tourism & Attractions": 55,
    "Tourism": 55,
    "Media, Marketing & Technology": 62,
    "Nonprofits, Charities & Community": 75,
    "Medical – Primary & Specialty Care": 78,
    "Education & Schools": 85,
    "Automotive": 80,
    "Telecom": 92,
    "Energy & Industrial": 95,
    "Banking & Credit Unions": 90,
    "Unclear / Needs Confirmation": 50,
}

KNOWN_INDUSTRY = {
    "208 companies": "Construction, Trades & Home Services",
    "a+ pioneer title": "Title & Escrow",
    "aci northwest": "Construction, Trades & Home Services",
    "adventure sport rentals": "Recreation, Tourism & Attractions",
    "alani bakery": "Restaurants, Food & Beverage",
    "allwest testing & engineering": "Engineering & Professional Services",
    "altek": "Engineering & Professional Services",
    "architects west, inc.": "Engineering & Professional Services",
    "architerra homes": "Construction, Trades & Home Services",
    "armon wealth management group of stifel": "Insurance & Financial Services",
    "atchley financial group": "Insurance & Financial Services",
    "avista": "Energy & Industrial",
    "avondale dental group": "Dentistry",
    "bales construction": "Construction, Trades & Home Services",
    "blue cross of idaho": "Insurance & Financial Services",
    "bouten construction": "Construction, Trades & Home Services",
    "caruso's sandwiches & artisan pizza": "Restaurants, Food & Beverage",
    "cda resort/hagadone hospitality": "Recreation, Tourism & Attractions",
    "coeur d'alene public schools: district 271": "Education & Schools",
    "coeur d'alene realty | windermere": "Real Estate Brokerages",
    "domino's": "Restaurants, Food & Beverage",
    "evan's brothers": "Restaurants, Food & Beverage",
    "fatbeam fiber": "Telecom",
    "fmi equipment": "Construction, Trades & Home Services",
    "fries construction": "Construction, Trades & Home Services",
    "gesa credit union": "Banking & Credit Unions",
    "hagadone marine group": "Recreation, Tourism & Attractions",
    "hagadone media group": "Media, Marketing & Technology",
    "harris & bruno international": "Energy & Industrial",
    "hecla mining company": "Energy & Industrial",
    "heritage health": "Medical – Primary & Specialty Care",
    "hmh engineering": "Engineering",
    "hospice of north idaho": "Home Care, Hospice & Senior Living",
    "hrei": "Nonprofits, Charities & Community",
    "idaho forest group": "Energy & Industrial",
    "itrip coeur d'alene": "Real Estate Brokerages",
    "j-u-b engineers, inc.": "Engineering",
    "kiemle hagood": "Real Estate Brokerages",
    "koerner furniture": "Retail & Consumer Goods",
    "kootenai county democrats": "Nonprofits, Charities & Community",
    "kootenai title": "Title & Escrow",
    "lakeside companies": "Insurance & Financial Services",
    "lewis-clark state college": "Education & Schools",
    "locate 208 real estate": "Real Estate Brokerages",
    "marcus anderson properties": "Apartments & Housing",
    "mcdonald's": "Restaurants, Food & Beverage",
    "mix it up": "Retail & Consumer Goods",
    "national mattress and furniture": "Retail & Consumer Goods",
    "north idaho college (nic)": "Education & Schools",
    "north idaho title": "Title & Escrow",
    "northwest bank": "Banking & Credit Unions",
    "oxyfresh": "Retail & Consumer Goods",
    "pacific office automation": "Engineering & Professional Services",
    "panhandle affordable housing alliance (paha)": "Nonprofits, Charities & Community",
    "paradigm of idaho": "Home Care, Hospice & Senior Living",
    "parker toyota": "Automotive",
    "pita pit usa": "Restaurants, Food & Beverage",
    "post falls er & hospital": "Medical – Primary & Specialty Care",
    "prairie family medicine": "Medical – Primary & Specialty Care",
    "rathdrum power": "Energy & Industrial",
    "regatta strategy group": "Media, Marketing & Technology",
    "regence blueshield of idaho": "Insurance & Financial Services",
    "republic services": "Energy & Industrial",
    "rojo ink": "Tattoo & Body Art",
    "springhill suites coeur d'alene": "Recreation, Tourism & Attractions",
    "st. vincent de paul north idaho": "Nonprofits, Charities & Community",
    "stcu": "Banking & Credit Unions",
    "super 1": "Retail & Consumer Goods",
    "tds telecom": "Telecom",
    "the business strategy institute": "Engineering & Professional Services",
    "the salvation army kroc center": "Nonprofits, Charities & Community",
    "tomlinson sothebys international realty": "Real Estate Brokerages",
    "university of idaho coeur d'alene": "Education & Schools",
    "victory media": "Media, Marketing & Technology",
    "watkins distributing": "Retail & Consumer Goods",
    "ziply fiber": "Telecom",
}

HARD_PATTERNS = [
    (r"\bus bank\b|\bwells fargo\b|\bfirst interstate\b|\bmountain west bank\b|\bnorthwest bank\b", 25),
    (r"\biccu\b|\bstcu\b|\bp1fcu\b|\bnumerica\b|\bhorizon credit\b|\bgoldenwest\b|\bgesa\b", 20),
    (r"\bverizon\b|\bziply\b|\btds telecom\b|\bfatbeam\b", 28),
    (r"\bavista\b|\brathdrum power\b|\bhecla\b|\brepublic services\b", 30),
    (r"\bblue cross\b|\bregence\b|\bnorthwestern mutual\b", 28),
    (r"\bhagadone\b|\bcda resort\b", 22),
    (r"\bschool district\b|\bpublic schools\b|\buniversity\b|\bcollege\b|north idaho college|\bnic\b", 22),
    (r"\bhospital\b|\ber &\b|\bnorthwest specialit|\bheritage health\b", 20),
    (r"mcdonald|domino|planet fitness|anytime fitness|snap fitness|club pilates|purebarre|stretchlab|drybar", 18),
    (r"\bidoaho forest group\b|\bbouten\b", 18),
    (r"windermere|sotheby|century 21", 8),
    (r"\bfoley financial\b|\bstifel\b", 15),
]

EASY_PATTERNS = [
    (r"\bsalon\b|\bnails\b|\bspa\b|\bmassage\b|\besthetic|beauty bar|aroma\b", -8),
    (r"\byoga\b|\bcrossfit\b|\bpersonal train|\bfitness\b|\bpilates\b", -3),
    (r"\bbakery\b|\bpizza\b|\bcoffee\b|\bchicken\b|\bwine\b|\bdistillery\b", -5),
    (r"\bfamily dental\b|\borthodont", -2),
    (r"\bdoula\b|\bnutrition\b|\bnaturopath|\bacupuncture|\bholistic", -6),
    (r"\btattoo\b|\bink\b|\bspray tan", -8),
    (r"\bcleaning\b|\bheating\b|\bfence\b|\bbackflow\b|\bexcavation\b", -4),
]


def normalize(name):
    if not name:
        return ""
    return re.sub(r"\s+", " ", str(name).replace("\xa0", " ").strip().lower())


def industry_score(industry):
    if not industry:
        return 55
    if str(industry).startswith("Medical"):
        return INDUSTRY_BASE["Medical – Primary & Specialty Care"]
    return INDUSTRY_BASE.get(industry, 55)


def company_adjust(name):
    n = normalize(name)
    adj = 0
    for pat, val in HARD_PATTERNS:
        if re.search(pat, n, re.I):
            adj += val
    for pat, val in EASY_PATTERNS:
        if re.search(pat, n, re.I):
            adj += val
    words = n.split()
    if len(words) <= 2 and not any(
        x in n for x in ["bank", "credit", "hospital", "university", "college", "verizon", "avista"]
    ):
        adj -= 3
    return adj


def lookup_known_industry(n):
    if n in KNOWN_INDUSTRY:
        return KNOWN_INDUSTRY[n]
    for k, v in KNOWN_INDUSTRY.items():
        if k == n or k in n or n in k:
            return v
    return None


def score_company(name, industry=None):
    n = normalize(name)
    ind = industry or lookup_known_industry(n)
    base = industry_score(ind)
    return base + company_adjust(name), ind or "Unknown"


def rank_sheet(ws, company_col, industry_col=None):
    company_meta = {}
    rows = []
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, company_col).value
        if not name or not str(name).strip():
            continue
        ind = ws.cell(r, industry_col).value if industry_col else None
        sc, inferred = score_company(name, ind)
        n = normalize(name)
        rows.append((r, n))
        if n not in company_meta or sc < company_meta[n][0]:
            company_meta[n] = (sc, str(name).strip(), inferred or ind)

    ordered = sorted(company_meta.items(), key=lambda x: (x[1][0], x[1][1].lower()))
    rank_map = {n: i + 1 for i, (n, _) in enumerate(ordered)}

    for r, n in rows:
        ws.cell(r, 1).value = rank_map[n]

    return ordered


print("Ranking Enriched...")
ordered1 = rank_sheet(wb["Enriched"], company_col=3, industry_col=2)
print(f"  {len(ordered1)} unique companies")
print("  Easiest 12:")
for i, (n, (sc, disp, ind)) in enumerate(ordered1[:12], 1):
    print(f"    {i}. {disp} | {ind} | score={sc}")
print("  Hardest 12:")
for i, (n, (sc, disp, ind)) in enumerate(ordered1[-12:], len(ordered1) - 11):
    print(f"    {i}. {disp} | {ind} | score={sc}")

print("\nRanking Needs Enriched...")
ordered2 = rank_sheet(wb["Needs Enriched"], company_col=3, industry_col=None)
print(f"  {len(ordered2)} unique companies")
print("  Easiest 12:")
for i, (n, (sc, disp, ind)) in enumerate(ordered2[:12], 1):
    print(f"    {i}. {disp} | {ind} | score={sc}")
print("  Hardest 12:")
for i, (n, (sc, disp, ind)) in enumerate(ordered2[-12:], len(ordered2) - 11):
    print(f"    {i}. {disp} | {ind} | score={sc}")

wb.save(path)
print("\nSaved:", path)
