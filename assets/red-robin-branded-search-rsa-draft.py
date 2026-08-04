# -*- coding: utf-8 -*-
"""Branded Search messaging Excel for Red Robin — 1:1 working file."""
import re
import subprocess
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    raise SystemExit("openpyxl required")


def effective_len(text: str) -> int:
    return len(re.sub(r"\{CUSTOMIZER\.[^:]+:([^}]+)\}", r"\1", text))


HEADER_FILL = PatternFill("solid", fgColor="8B1E1E")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(bold=True, size=14)
SECTION_FONT = Font(bold=True, size=11)
THIN = Border(
    left=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
    top=Side(style="thin", color="D0D0D0"),
    bottom=Side(style="thin", color="D0D0D0"),
)
WRAP = Alignment(wrap_text=True, vertical="top")


def style_header_row(ws, row: int) -> None:
    for cell in ws[row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")


def autosize(ws, widths: dict[str, int]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


# ---------------------------------------------------------------------------
# Copy banks — recurring + family / kids + Royalty + catering
# Sourced from redrobin.com/rewards, /menu/kids-meals, /catering
# ---------------------------------------------------------------------------

# Territory B lead (recommended): Belong & Reward for returning guests
BRAND_HEADLINES = [
    ("H1", "Returning guest", "Welcome Back to Red Robin®", "Recurring searcher; brand already known"),
    ("H2", "Loyalty CTA", "Join Red Robin Royalty®", "https://www.redrobin.com/rewards"),
    ("H3", "Signup perk", "Free Appetizer for Royalty*", "Sign up perk; $10 min; terms apply"),
    ("H4", "Points math", "Earn $10 for 100 Points", "1 pt / $1 → $10 at 100"),
    ("H5", "Birthday", "Birthday Treats for Royalty", "Birthday month perk"),
    ("H6", "Member exclusives", "Member-Only Deals Inside", "Exclusive member offers"),
    ("H7", "Family lead", "Family Dining Done Right", "Core brand-intent family angle"),
    ("H8", "Kids invite", "Kids Love Red Robin®", "From kids meals page"),
    ("H9", "Kids menu", "Kids Meals & Bottomless Fries", "Kids entrées + Bottomless Steak Fries®"),
    ("H10", "Kids faves", "Cluck-A-Doodles & Mac It Yours", "Named kids favorites from menu"),
    ("H11", "Whole crew", "Bring the Whole Family", "Table-time / group dining"),
    ("H12", "Bottomless habit", "Bottomless Sides & Drinks", "Signature — not promo-led"),
    ("H13", "Catering CTA", "Catering for Any Occasion", "https://www.redrobin.com/catering"),
    ("H14", "Catering reward", "Earn Catering Rewards", "$25 eGift per $500 spent"),
    ("H15", "Order path", "Order Pickup or Delivery", "App / online path for returners"),
]

ALT_HEADLINES = [
    ("Alt", "App ordering", "Easier Ordering in the App", "Rewards page app CTA"),
    ("Alt", "Gourmet burger bar", "Host a Gourmet Burger Bar", "Catering signature offer"),
    ("Alt", "Boxed meals", "Boxed Meals for Your Crew", "Catering boxed meals"),
    ("Alt", "Kids bundles", "Kids & Pizza Catering Bundles", "Catering bundles"),
    ("Alt", "Kids drinks", "Kids Dirty Sodas & Lemonades", "Kids beverage menu"),
    ("Alt", "Early access", "Early Access for Royalty", "Member sneak peeks"),
    ("Alt", "Points every visit", "Earn Points Every Visit", "Recurring-guest habit"),
    ("Alt", "Local CTA", "Visit Your Local Red Robin®", "Keep light local line"),
    ("Alt", "Soft offer garnish", "{CUSTOMIZER.Deal:Big Yummm Deals}", "Territory C garnish only"),
    ("Alt", "Soft price garnish", "From {CUSTOMIZER.DealPrice:$9.99}", "1 price pin max on Brand"),
]

DESCRIPTIONS = [
    (
        "D1",
        "Royalty + returning",
        "Back for more? Join Royalty® for a free appetizer*, points every visit & member deals.",
        "Recurring + loyalty lead",
    ),
    (
        "D2",
        "Family + kids",
        "Bring the crew. Kids meals with Bottomless Steak Fries®—Cluck-A-Doodles & Mac It Yours.",
        "Family restaurant proof from kids menu",
    ),
    (
        "D3",
        "Royalty math",
        "Earn 1 point per $1—100 points = $10 reward. Unlock bottomless Royalty® perks today.",
        "Rewards page math",
    ),
    (
        "D4",
        "Catering",
        "Feed the team or party. Gourmet Burger Bar, boxed meals & bundles—earn catering rewards.",
        "From catering page",
    ),
]

DESC_BANK = [
    (
        "D5",
        "Family + Royalty",
        "Family table time starts here. Earn Royalty® points every visit—kids favorites included.",
        "Recurring family combo",
    ),
    (
        "D6",
        "Soft deal garnish",
        "While you’re here, see {CUSTOMIZER.Deal:Big Yummm Deals} from {CUSTOMIZER.DealPrice:$9.99}—plus Royalty® rewards.",
        "Optional C garnish; not the Brand lead",
    ),
]

SITELINKS = [
    (
        "Red Robin Royalty®",
        "Free appetizer* & points every visit.",
        "Birthday treats & member-only deals.",
        "https://www.redrobin.com/rewards",
    ),
    (
        "Kids Meals",
        "Cluck-A-Doodles, Mac It Yours & more.",
        "Bottomless Steak Fries® with kids meals.",
        "https://www.redrobin.com/menu/kids-meals",
    ),
    (
        "Catering",
        "Gourmet Burger Bar, boxes & bundles.",
        "Earn $25 eGift per $500 catering.",
        "https://www.redrobin.com/catering",
    ),
    (
        "Bottomless Menu",
        "30+ bottomless sides & drinks.",
        "Unlimited refills on the good stuff.",
        "https://www.redrobin.com/menu/bottomless",
    ),
    (
        "Big Yummm Deals",
        "Full meal deals starting at $9.99.",
        "Entrée, bottomless side & beverage.",
        "https://www.redrobin.com/the-big-yummm",
    ),
    (
        "Gift Cards",
        "Say it with burgers and brews.",
        "eGift or physical—delivered fast.",
        "https://www.redrobin.com/gift-cards",
    ),
]

SOURCE_NOTES = [
    (
        "Rewards",
        "https://www.redrobin.com/rewards",
        "Royalty®: free appetizer* ($10 min), 1 pt/$1, 100 pts=$10, birthday treat, member exclusives, app ordering (curbside/pickup/delivery)",
    ),
    (
        "Kids Meals",
        "https://www.redrobin.com/menu/kids-meals",
        "Kids Love Red Robin; Red’s Cheeseburger, Mac It Yours (bottomless), Cluck-A-Doodles, Corn Doggies, Grilled Chicken Dip’ns, Dirty Sodas, Bottomless Steak Fries®",
    ),
    (
        "Catering",
        "https://www.redrobin.com/catering",
        "Gourmet Burger Bar, boxed meals, catering bundles (kids/pizza/wings/burgers); $25 eGift per $500 catering spend; call 833-224-5340",
    ),
]


def write_brief(wb) -> None:
    ws = wb.create_sheet("1 Brief and Stance", 0)
    ws["A1"] = "Red Robin · Branded Search Messaging"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:B1")

    rows = [
        [],
        ["Lens", "Searcher already Googled “Red Robin.” They know us and chose us — often a returning / recurring guest."],
        ["Problem today", "Brand RSAs re-introduce the brand and restate $9.99 BYD — same copy as Competitor & Menu."],
        ["Inventory", "15 RSAs · 6 brand campaigns · 2 variants (BYD + Kids) · 3 Poor / 7 Average ad strength"],
        [],
        ["WORKING ANSWERS"],
        [
            "Job of Brand ads",
            "Not awareness. Win paid slot vs brand-bidders, route next action, deepen relationship for returners. Reassurance/offer support — don’t lead.",
        ],
        [
            "BYD / $9.99",
            "Not the lead on Exact/Phrase Brand. At most 1–2 customizer lines or sitelink. Full BYD stays on Competitor + Menu.",
        ],
        [
            "Recurring guests",
            "Lead with Royalty (points, free app*, birthday, member deals) + family table habit — not “who we are.”",
        ],
        [
            "Kids / family",
            "Brand should carry light kids proof (Kids Love Red Robin, named faves, Bottomless fries). Full Kids RSA stays on Menu - Kids.",
        ],
        [
            "Catering",
            "Sitelink + 1–2 headlines for returners who feed groups (Burger Bar, catering rewards). Don’t make Brand a catering campaign.",
        ],
        [
            "Own Brand RSA?",
            "Yes — stop sharing BYD with Competitor/Menu. Kids already has its own set.",
        ],
        [
            "Geo / Near Me",
            "Proximity = targeting + location extensions. Same creative as Brand Exact/Phrase.",
        ],
        [],
        ["LEAN FOR 1:1", "Territory B (Belong & Reward) + family/kids + Royalty. Optional soft BYD garnish. Competitor/Menu keep full BYD."],
    ]

    r = 2
    for row in rows:
        for c, val in enumerate(row, 1):
            cell = ws.cell(r, c, val)
            cell.alignment = WRAP
            if len(row) == 1 and "WORKING" in str(row[0]):
                cell.font = SECTION_FONT
            if c == 1 and len(row) == 2:
                cell.font = Font(bold=True)
        r += 1

    autosize(ws, {"A": 22, "B": 100})
    ws.row_dimensions[3].height = 35
    for i in range(7, 15):
        ws.row_dimensions[i].height = 40


def write_territories(wb) -> None:
    ws = wb.create_sheet("2 Territories")
    ws["A1"] = "Three messaging territories (sample headlines — not finished RSA)"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:E1")

    ws.append([])
    headers = ["Territory", "Job / when it wins", "Sample headlines", "BYD role", "Best for"]
    ws.append(headers)
    style_header_row(ws, ws.max_row)

    territories = [
        [
            "A — Route & Convert",
            "Click is nearly free; reduce friction: locator, order, hours, dine-in vs pickup.\nWins when KPI = orders / store visits.",
            "Start Your Order Online\nFind Your Red Robin\nPickup, Delivery or Dine-In\nSee Hours & Directions\nOrder Ahead Tonight\nMenu Ready When You Are",
            "Optional 1 supporting line or sitelink only — not the lead.",
            "Brand Exact/Phrase if KPI = orders; Geo inherits same set.",
        ],
        [
            "B — Belong & Reward  [LEAN]",
            "They already chose RR (often returning) — deepen relationship: Royalty, family table, kids, bottomless habit.\nWins when KPI = loyalty / frequency / LTV.",
            "Welcome Back to Red Robin®\nJoin Red Robin Royalty®\nFree Appetizer for Royalty*\nEarn $10 for 100 Points\nKids Love Red Robin®\nKids Meals & Bottomless Fries\nBring the Whole Family\nCatering for Any Occasion\nBottomless Sides Await",
            "Absent from lead headlines; sitelink or soft description mention only.",
            "Brand Exact/Phrase; strongest for recurring / family guests.",
        ],
        [
            "C — Offer as Accelerator",
            "Brand intent + value nudge: deal is present but secondary — “while you’re here.”\nWins as garnish if leadership won’t fully drop BYD from Brand.",
            "Your Table Is Waiting\nTonight’s Deal Is On\n{Deal} While You’re Here\nFrom {Price} Full Meals\nBottomless Still Included\nCome Hungry. Leave Happy.",
            "1–2 customizer headlines + 1 description max. Not 8+ offer lines like today’s BYD set.",
            "Optional hybrid with B; bridge if BYD must stay visible on Brand.",
        ],
    ]
    for t in territories:
        ws.append(t)
        for cell in ws[ws.max_row]:
            cell.alignment = WRAP
            cell.border = THIN
        ws.row_dimensions[ws.max_row].height = 130

    ws.append([])
    ws.append(["Variant architecture"])
    ws[f"A{ws.max_row}"].font = SECTION_FONT
    ws.append(["Ad group type", "RSA variant", "Lead message", "Notes"])
    style_header_row(ws, ws.max_row)
    for row in [
        ["Brand Exact / Phrase", "NEW — Brand (B, or B+C garnish)", "Royalty + family/kids; offer secondary", "Stop sharing BYD set"],
        ["Geo / Near Me / Locations", "Same as Brand Exact/Phrase", "Same creative; location via targeting + extensions", "Proximity ≠ headline job"],
        ["Competitor", "Keep BYD (offer-led)", "$9.99 / Big Yummm / vs alt", "Still deciding vs alternatives"],
        ["Menu (Deals / General / Category)", "Keep BYD; Kids stays Kids", "Offer + menu proof", "Menu-Deals especially offer-led"],
        ["Menu - Kids", "Keep Kids variant", "Kids menu specifics", "Already differentiated — Brand only borrows light kids lines"],
    ]:
        ws.append(row)
        for cell in ws[ws.max_row]:
            cell.alignment = WRAP

    autosize(ws, {"A": 32, "B": 48, "C": 42, "D": 40, "E": 42})


def write_rsa(wb) -> None:
    ws = wb.create_sheet("3 Brand RSA B lean")
    ws["A1"] = "Brand Exact/Phrase RSA draft — Territory B lean (recurring + family/kids + Royalty)"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:G1")
    ws["A2"] = (
        "Sample / working draft for 1:1. Rebuild pin preferences after direction lock. "
        "Legal: free appetizer* and rewards claims need disclosure where required."
    )
    ws["A2"].alignment = WRAP
    ws.merge_cells("A2:G2")
    ws.row_dimensions[2].height = 35

    ws.append([])
    ws.append(["HEADLINES (30-char limit)"])
    ws[f"A{ws.max_row}"].font = SECTION_FONT
    ws.append(["Slot", "Theme", "Copy", "Chars", "Limit", "Headroom", "Notes / source"])
    style_header_row(ws, ws.max_row)
    for slot, theme, copy, notes in BRAND_HEADLINES:
        n = effective_len(copy)
        ws.append([slot, theme, copy, n, 30, 30 - n, notes])
        for cell in ws[ws.max_row]:
            cell.alignment = WRAP
            if isinstance(cell.value, int) and cell.column == 6 and cell.value < 0:
                cell.font = Font(color="C00000", bold=True)

    ws.append([])
    ws.append(["ALTERNATE HEADLINES (swap bank)"])
    ws[f"A{ws.max_row}"].font = SECTION_FONT
    ws.append(["Slot", "Theme", "Copy", "Chars", "Limit", "Headroom", "Notes / source"])
    style_header_row(ws, ws.max_row)
    for slot, theme, copy, notes in ALT_HEADLINES:
        n = effective_len(copy)
        ws.append([slot, theme, copy, n, 30, 30 - n, notes])
        for cell in ws[ws.max_row]:
            cell.alignment = WRAP

    ws.append([])
    ws.append(["DESCRIPTIONS (90-char limit)"])
    ws[f"A{ws.max_row}"].font = SECTION_FONT
    ws.append(["Slot", "Theme", "Copy", "Chars", "Limit", "Headroom", "Notes / source"])
    style_header_row(ws, ws.max_row)
    for slot, theme, copy, notes in DESCRIPTIONS + DESC_BANK:
        n = effective_len(copy)
        ws.append([slot, theme, copy, n, 90, 90 - n, notes])
        for cell in ws[ws.max_row]:
            cell.alignment = WRAP

    autosize(ws, {"A": 8, "B": 20, "C": 55, "D": 8, "E": 8, "F": 10, "G": 48})


def write_sources(wb) -> None:
    ws = wb.create_sheet("4 Source copy")
    ws["A1"] = "Live site proof points used in Brand copy"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:C1")

    ws.append([])
    ws.append(["Page", "URL", "Pullable claims / language"])
    style_header_row(ws, ws.max_row)
    for row in SOURCE_NOTES:
        ws.append(list(row))
        for cell in ws[ws.max_row]:
            cell.alignment = WRAP
        ws.row_dimensions[ws.max_row].height = 55

    ws.append([])
    ws.append(["How this maps to recurring guests"])
    ws[f"A{ws.max_row}"].font = SECTION_FONT
    for line in [
        "Returning searchers don’t need “Red Robin’s Big Yummm Deals” — they need a reason to come back now or deepen the habit.",
        "Royalty = the returning-guest product: points every visit, free app*, birthday, member deals, app reorder.",
        "Kids = why families keep choosing RR (named kids faves + Bottomless fries) without turning Brand into the Kids RSA.",
        "Catering = upsell for guests who already trust the brand and feed groups — headline + sitelink, not the whole ad.",
    ]:
        ws.append([line])
        ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=3)
        ws[f"A{ws.max_row}"].alignment = WRAP
        ws.row_dimensions[ws.max_row].height = 32

    autosize(ws, {"A": 14, "B": 42, "C": 95})


def write_sitelinks(wb) -> None:
    ws = wb.create_sheet("5 Sitelinks")
    ws["A1"] = "Recommended branded sitelinks (pair with Territory B)"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:D1")

    ws.append([])
    ws.append(["Link text", "Description 1", "Description 2", "Final URL"])
    style_header_row(ws, ws.max_row)
    for row in SITELINKS:
        ws.append(list(row))
        for cell in ws[ws.max_row]:
            cell.alignment = WRAP

    autosize(ws, {"A": 24, "B": 48, "C": 48, "D": 48})


def write_checklist(wb) -> None:
    ws = wb.create_sheet("6 One-on-one checklist")
    ws["A1"] = "Decision checklist"
    ws["A1"].font = TITLE_FONT

    ws.append([])
    for item in [
        "1. Primary Brand KPI: store visit / order / Royalty signup / defend SERP?",
        "2. Lock Territory B (or A / C / B+C garnish).",
        "3. Confirm Brand Exact/Phrase gets its own RSA; Competitor + Menu keep BYD.",
        "4. Confirm light kids lines on Brand OK — full Kids RSA stays on Menu - Kids.",
        "5. Confirm catering = sitelink + 1–2 H, not a Brand theme takeover.",
        "6. Confirm Geo / Near Me clones Brand creative (not a third “near me” set).",
        "7. After pick: build/pin full 15H/4D + sitelinks; retire shared BYD on Brand AGs.",
        "8. Watch ad strength on the 3 Poor / 7 Average RSAs after split.",
    ]:
        ws.append([item])
        ws[f"A{ws.max_row}"].alignment = WRAP
        ws.row_dimensions[ws.max_row].height = 22

    autosize(ws, {"A": 110})


def main() -> None:
    out = Path(r"C:\Users\skyes\OneDrive\Desktop\Red Robin Branded Search Messaging.xlsx")
    # Fallback if OneDrive Desktop path differs
    if not out.parent.exists():
        out = Path.home() / "Desktop" / "Red Robin Branded Search Messaging.xlsx"
    # If first open locked the file, write a refreshed copy beside it
    try:
        test = open(out, "a")
        test.close()
    except PermissionError:
        out = out.with_name("Red Robin Branded Search Messaging v2.xlsx")

    wb = openpyxl.Workbook()
    default = wb.active
    wb.remove(default)

    write_brief(wb)
    write_territories(wb)
    write_rsa(wb)
    write_sources(wb)
    write_sitelinks(wb)
    write_checklist(wb)

    print("Char QA — Brand headlines")
    for slot, theme, copy, _ in BRAND_HEADLINES:
        n = effective_len(copy)
        flag = "OK" if n <= 30 else "OVER"
        print(f"  {slot} {n:2d}/30 {flag} | {copy}")
    print("Char QA — Descriptions")
    for slot, theme, copy, _ in DESCRIPTIONS + DESC_BANK:
        n = effective_len(copy)
        flag = "OK" if n <= 90 else "OVER"
        print(f"  {slot} {n:2d}/90 {flag} | {copy}")

    wb.save(out)
    print(f"\nWrote {out}")

    # Open on Windows
    subprocess.Popen(["cmd", "/c", "start", "", str(out)], shell=False)
    print("Opened in default Excel app.")


if __name__ == "__main__":
    main()
